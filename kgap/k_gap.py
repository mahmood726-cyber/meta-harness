"""K-GAP measurement: the comparator's INCLUDED-STUDIES set vs our pool, trial by trial.

A comparator's headline k is a number; its included-studies table is a SET. Parity on the number
hides which trials are missing and why. This module reads the comparator's own JATS (PMC OA) where
table cells cite the reference list by `<xref ref-type="bibr" rid=...>`, and the reference list
carries `<pub-id pub-id-type="pmid">` -- so the comparator's trial set is extracted by PARSING, with
no model and no prose regex over trial names. Every row keeps the table label, the verbatim cell text
and the rid it cited, so each membership claim has a span.

Fetch (network, once, cached with sha256 under cache/comparators/<pmid>/) is separated from the
parse/classify steps, which are pure functions of cached bytes.

Nothing here admits a trial or changes a pool. The output is a MEASUREMENT: which comparator trials we
lack, and which open source could supply an admissible, typed result for each.
"""
from __future__ import annotations

import hashlib
import json
import os
import re
import xml.etree.ElementTree as ET
from typing import Any

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
COMP_DIR = os.path.join(ROOT, "cache", "comparators")

_WS = re.compile(r"\s+")
NCT_RE = re.compile(r"\bNCT\d{8}\b")

# A table is an included-studies table when its caption says so. Deliberately narrow: an outcome table
# ("Subgroup analysis of ...") also cites trials, but its rows are analyses, not the included set.
INCLUDED_CAPTION = re.compile(
    r"characteristic|included (?:stud|trial|rct)|baseline|summary of (?:the )?(?:included )?(?:stud|trial)"
    r"|(?:stud|trial)s? included|description of (?:the )?(?:included )?(?:stud|trial)|study design|overview of|summary of (?:the )?(?:cvot|outcome trial|rct|randomi)",
    re.I)


def _flat(s: str | None) -> str:
    return _WS.sub(" ", s or "").strip()


def _text(el) -> str:
    return _flat("".join(el.itertext())) if el is not None else ""


def sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


# ----------------------------------------------------------------------------------------------- fetch

def _cache_write(pmid: str, name: str, body: bytes) -> dict:
    d = os.path.join(COMP_DIR, pmid)
    os.makedirs(d, exist_ok=True)
    p = os.path.join(d, name)
    with open(p, "wb") as fh:
        fh.write(body)
    return {"file": os.path.relpath(p, ROOT).replace(os.sep, "/"), "bytes": len(body), "sha256": sha256(body)}


def fetch_comparator_jats(pmid: str, date: str) -> dict:
    """idconv PMID->PMCID, then Europe PMC fullTextXML, falling back to NCBI efetch db=pmc.

    Returns a manifest dict; state is JATS_FULL (has <body> or table-wrap and ref-list), JATS_FRONT_ONLY
    (publisher withholds the body), NO_PMCID, or FETCH_FAILED. Never raises on a missing source."""
    from harness import http
    man: dict[str, Any] = {"pmid": pmid, "date": date, "records": []}
    name = f"{date}_kgap_jats.xml"
    existing = os.path.join(COMP_DIR, pmid, name)
    if os.path.exists(existing):
        body = open(existing, "rb").read()
        man.update(_classify_jats(body))
        man["records"].append({"state": "CACHED", **_rel(existing, body)})
        return man
    try:
        st, b = http.get_raw("https://www.ncbi.nlm.nih.gov/pmc/utils/idconv/v1.0/",
                             {"ids": pmid, "format": "json", "tool": "meta-harness", "email": "meta-harness@example.org"})
        rec = _cache_write(pmid, f"{date}_kgap_idconv.json", b)
        man["records"].append({"url": "idconv", "http_status": st, **rec})
        recs = json.loads(b.decode("utf-8")).get("records") or [{}]
        pmcid = recs[0].get("pmcid") or ""
    except Exception as exc:  # noqa: BLE001
        man.update({"state": "FETCH_FAILED", "error": f"idconv: {exc}"[:300]})
        return man
    man["pmcid"] = pmcid
    if not pmcid:
        man["state"] = "NO_PMCID"
        return man
    for url, params in (
        (f"https://www.ebi.ac.uk/europepmc/webservices/rest/{pmcid}/fullTextXML", None),
        ("https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi", {"db": "pmc", "id": pmcid, "retmode": "xml"}),
    ):
        try:
            st, b = http.get_raw(url, params, tries=2)
        except Exception as exc:  # noqa: BLE001
            man["records"].append({"url": url, "error": str(exc)[:200]})
            continue
        cls = _classify_jats(b)
        if cls["state"] == "JATS_FULL":
            rec = _cache_write(pmid, name, b)
            man["records"].append({"url": url, "http_status": st, **rec})
            man.update(cls)
            return man
        man["records"].append({"url": url, "http_status": st, "state": cls["state"], "bytes": len(b)})
    man["state"] = "JATS_FRONT_ONLY"
    return man


def _rel(p, body):
    return {"file": os.path.relpath(p, ROOT).replace(os.sep, "/"), "bytes": len(body), "sha256": sha256(body)}


def _classify_jats(body: bytes) -> dict:
    try:
        root = ET.fromstring(body)
    except ET.ParseError:
        return {"state": "NOT_XML"}
    has_refs = root.find(".//ref-list") is not None
    has_tab = root.find(".//table-wrap") is not None
    has_body = root.find(".//body") is not None
    return {"state": "JATS_FULL" if (has_refs and (has_tab or has_body)) else "JATS_FRONT_ONLY",
            "has_tables": has_tab, "has_refs": has_refs}


# ----------------------------------------------------------------------------------------------- parse

def parse_refs(root) -> dict[str, dict]:
    out = {}
    for ref in root.iter("ref"):
        rid = ref.get("id")
        if not rid:
            continue
        ids = {}
        for pid in ref.iter("pub-id"):
            t = (pid.get("pub-id-type") or "").lower()
            if t in ("pmid", "doi", "pmcid") and pid.text:
                ids[t] = pid.text.strip()
        txt = _text(ref)
        if "doi" not in ids:
            m = re.search(r"\b(10\.\d{4,9}/[^\s\"<>]+[^\s\"<>.,;])", txt)
            if m:
                ids["doi"] = m.group(1)
        title = ""
        at = ref.find(".//article-title")
        if at is not None:
            title = _text(at)
        year = ""
        y = ref.find(".//year")
        if y is not None and y.text:
            year = y.text.strip()
        surname = ""
        sn = ref.find(".//surname")
        if sn is not None and sn.text:
            surname = sn.text.strip()
        lab = ref.find("label")
        ref_label = _text(lab).strip(" .[]()") if lab is not None else ""
        out[rid] = {"rid": rid, **ids, "title": title, "year": year, "first_author": surname,
                    "label": ref_label, "ordinal": len(out) + 1,
                    "text": txt[:400], "ncts": sorted(set(NCT_RE.findall(txt)))}
    return out


def _cell_rids(cell) -> list[str]:
    rids = []
    for x in cell.iter("xref"):
        if (x.get("ref-type") or "") == "bibr":
            for r in (x.get("rid") or "").split():
                rids.append(r)
    return rids


def parse_tables(root) -> list[dict]:
    tabs = []
    for tw in root.iter("table-wrap"):
        label = _text(tw.find("label"))
        caption = _text(tw.find("caption"))
        header = []
        rows = []
        for tr in tw.iter("tr"):
            cells = [c for c in tr if c.tag in ("td", "th")]
            if not cells:
                continue
            texts = [_text(c) for c in cells]
            if all(c.tag == "th" for c in cells) and not rows:
                # keep each header cell's citations: a TRANSPOSED table (trials as columns, e.g. the
                # tranexamic comparator's "WOMAN | WOMAN-2 | TRAAP ...") cites its trials only here.
                header.append([{"text": t, "rids": _cell_rids(c)} for t, c in zip(texts, cells)])
                continue
            rids = []
            for c in cells:
                rids.extend(_cell_rids(c))
            rows.append({"cells": texts, "rids": rids, "row_text": " | ".join(texts)})
        tabs.append({"label": label, "caption": caption, "header": header, "rows": rows,
                     "is_included_table": bool(INCLUDED_CAPTION.search(caption or "") or INCLUDED_CAPTION.search(label))})
    return tabs


def parse_jats(body: bytes) -> dict:
    root = ET.fromstring(body)
    title = _text(root.find(".//article-meta//article-title"))
    return {"title": title, "refs": parse_refs(root), "tables": parse_tables(root)}


# ------------------------------------------------------------------------------ included-set extraction

# A trial label is a row's first cell. Labels that are table furniture (a sub-header, a units row, a
# "Year" continuation row) are not trials; they carry neither a citation, an NCT, nor a name token.
_FURNITURE = re.compile(r"^(?:study|trial|author|year|\d{4}|n|total|overall|reference|characteristics?|"
                        r"participants?|treatment|control|intervention|placebo|mean|median|[-–— ]*)$", re.I)
# a group / arm / stratum descriptor carrying its own n ('Stain used group in QRISK 10–19% (n = 6438)')
_STRATUM_ROW = re.compile(r"\b(?:group|arm|stratum|strata|subgroup|cohort)\b.*\(\s*n\s*=\s*[\d,]+\s*\)\s*$", re.I)
# after a HYPHEN a segment may be mixed case ('EMPEROR-Reduced', 'EMPEROR-Preserved'); after a SPACE only caps/digits
# ('PIONEER 6', 'ENGAGE AF-TIMI 48'), so an ordinary word ('RALES Study') is not absorbed into the acronym
_ACRO = re.compile(r"\b([A-Z][A-Z0-9]{2,}(?:-[A-Z0-9][A-Za-z0-9]*| [A-Z0-9]{1,}\b){0,3})")
# a surname may carry particles ('van der Molen', 'de Vrese') or be two capitalised words ('Ben Ayed', 'Almeida
# Montes'): those labels were read as having NO author-year key and the trial was never searched
_SURNAME = r"(?:(?:[Vv]an|[Dd]e[rn]?|[Vv]on|[Dd][aiu]|[Dd]el|[Dd]os|[Ll][ae]|[Ee]l|[Aa]l)\s+){0,2}[A-Z][A-Za-z'À-ſ-]+(?:\s+[A-Z][A-Za-z'À-ſ-]{2,})?"
_AUTHOR_YEAR = re.compile(r"^(" + _SURNAME + r")(?:\s+et\s+al\.?)?,?\s*\(?((?:19|20)\d\d)\)?")
_PAREN_ACRO = re.compile(r"\(([A-Z][A-Za-z0-9]*[A-Z0-9][A-Za-z0-9]*(?:[- ][A-Za-z0-9]+){0,3})\)")
_NOT_ACRO = {"RCT", "RCTS", "USA", "UK", "NA", "NR", "HR", "RR", "OR", "CI", "BMI", "LDL", "HDL", "CKD", "HF",
             "HFREF", "HFPEF", "LVEF", "NYHA", "ACS", "MI", "CAD", "PCI", "CABG", "DM", "T2DM", "AF", "VTE",
             "DVT", "PE", "ITT", "MACE", "COVID", "SARS", "ICU", "ARDS", "CAP", "PPH", "MADRS", "TRD",
             "EGFR", "UACR", "SBP", "DBP", "BID", "QD", "IV", "PO", "SC", "EPA", "DHA", "IU", "TID", "NCT",
             "DOAC", "DOACS", "VKA", "NOAC", "SGLT2", "SGLT2I", "GLP", "GLP1", "RA", "DPP4", "PCSK9", "FCM",
             "IPD", "HFPEF", "LCZ696", "ARNI", "ACEI", "ARB", "MRA", "MRAS", "NR", "NS", "YES", "NO", "ALL"}


_DASHES = re.compile("[‐-―−﹣－]")


def fold_dashes(s: str) -> str:
    """Typographic dashes -> '-'. 'DAPA‐HF' (U+2010) was tokenised as 'DAPA' and resolved to an unrelated
    DAPA-named registration: a WRONG identity, not an unresolved one."""
    return _DASHES.sub("-", s or "")


_APOSTROPHES = str.maketrans({"’": "'", "‘": "'", "ʼ": "'", "′": "'"})


def _label_tokens(label: str) -> dict:
    # typographic apostrophes -> "'": 'O’Neil, 2018' read no author, so the row was skipped as furniture
    lab = fold_dashes(_flat(label)).translate(_APOSTROPHES)
    acr = [a for a in _ACRO.findall(lab) if a.replace("-", "").replace(" ", "").upper() not in _NOT_ACRO
           and not NCT_RE.match(a)]
    m = _AUTHOR_YEAR.match(lab)
    return {"acronyms": acr[:3], "author": m.group(1) if m else "", "year": m.group(2) if m else "",
            "ncts": sorted(set(NCT_RE.findall(lab)))}


def _units_from_table(t: dict) -> list[dict]:
    """A table's trial units: its ROWS normally; its header COLUMNS when the header cites trials
    (transposed layout). A row with no citation still counts when its first cell carries an NCT, an
    acronym or an Author-Year label -- that is a trial named without a reference link."""
    units = []
    hdr = t["header"][0] if t["header"] else []
    hdr_cited = [h for h in hdr[1:] if h["rids"]]
    if len(hdr_cited) >= 2 or (hdr and not any(r["rids"] for r in t["rows"]) and len(hdr) >= 3
                               and all(_label_tokens(h["text"])["acronyms"] for h in hdr[1:] if h["text"])):
        for i, h in enumerate(hdr[1:], start=1):
            if not h["text"] and not h["rids"]:
                continue
            col = " | ".join(r["cells"][i] for r in t["rows"] if len(r["cells"]) > i)
            units.append({"layout": "column", "label": h["text"], "rids": h["rids"], "context": col[:600]})
        return units
    for r in t["rows"]:
        first = r["cells"][0] if r["cells"] else ""
        toks = _label_tokens(first)
        if not r["rids"] and not toks["ncts"] and not toks["acronyms"] and not toks["author"]:
            continue
        if _FURNITURE.match(_flat(first)) and not r["rids"]:
            continue
        if not r["rids"] and not toks["ncts"] and not any(_flat(c).strip() for c in r["cells"][1:]):
            # a SECTION HEADER spanning the table ('GLP-1 RA vs. placebo' above its trials): a label and nothing else,
            # no citation, no NCT -- not a trial ('GLP-1 RA' even reads as an acronym the stop-list does not hold)
            continue
        if not r["rids"] and not toks["ncts"] and _STRATUM_ROW.search(_flat(first)):
            # an ARM / STRATUM row of the study above it ('No stain used group in QRISK <10% (n = 39866)': 15 QRISK strata
            # of one cohort study, statins-primary-prevention-elderly 39076238) -- a group with its own n, no citation,
            # no NCT -- is a sub-row of a unit, never a unit
            continue
        units.append({"layout": "row", "label": first, "rids": r["rids"], "context": r["row_text"][:600]})
    return units


def included_trials(parsed: dict, agent_terms: list[str], other_agents: list[str] | None = None) -> dict:
    """Units (rows, or columns of a transposed table) of the comparator's included-studies table(s).

    Each unit keeps the verbatim label, the table it came from, the reference-list entries it cites
    (pmid/doi from the JATS ref-list), NCTs written in it, and acronym / Author-Year tokens for the
    identity resolver. agent_terms: the topic's agent names. A unit is DRUG_MATCH when its text names
    one; when NO unit names any topic agent the table is single-agent (the agent is in the paper's title,
    not a column) and units are AGENT_IMPLICIT -- recorded as such, never silently promoted.
    """
    refs = parsed["refs"]
    agent_re = re.compile("|".join(re.escape(a) for a in agent_terms), re.I) if agent_terms else None
    units, used, seen = [], [], set()
    for t in parsed["tables"]:
        if not t["is_included_table"]:
            continue
        tu = _units_from_table(t)
        if not tu:
            continue
        used.append(t["label"] or t["caption"][:60])
        for u in tu:
            toks = _label_tokens(u["label"])
            key = tuple(sorted(set(u["rids"]))) or ("L", _flat(u["label"]).upper())
            if key in seen:
                continue
            seen.add(key)
            cited = [refs[x] for x in dict.fromkeys(u["rids"]) if x in refs]
            text = u["label"] + " | " + u["context"]
            units.append({
                "table": t["label"] or t["caption"][:40],
                "layout": u["layout"],
                "label": _flat(u["label"])[:160],
                "context": _flat(u["context"])[:600],
                "rids": list(dict.fromkeys(u["rids"])),
                "cited": [{k: c.get(k) for k in ("rid", "pmid", "doi", "pmcid", "title", "year", "first_author", "ncts")}
                          for c in cited],
                "ncts": sorted(set(NCT_RE.findall(text)) | {n for c in cited for n in c.get("ncts", [])}),
                "acronyms": toks["acronyms"], "author": toks["author"], "year": toks["year"],
                "agent_hit": bool(agent_re and agent_re.search(text)),
            })
    other_re = re.compile("|".join(re.escape(a) for a in other_agents), re.I) if other_agents else None
    for u in units:
        # A unit is OTHER_AGENT only on its OWN text: it names another served topic's agent and none of ours.
        # (An earlier rule -- "some other row names our agent, so this one must not be ours" -- discarded 8 of 9
        # colchicine rows and the WOMAN/TRAAP columns: absence of a drug name in a row is not another drug.)
        text = u["label"] + " | " + u["context"]
        if u["agent_hit"]:
            u["drug_match"] = "DRUG_MATCH"
        elif other_re and other_re.search(text):
            u["drug_match"] = "OTHER_AGENT"
        else:
            u["drug_match"] = "AGENT_IMPLICIT"
    return {"tables_used": used, "units": units,
            "state": "ENUMERATED" if units else "NO_INCLUDED_TABLE_ENUMERABLE"}


# ------------------------------------------------------------------------------------ AACT local index

def norm_acronym(a: str) -> str:
    """'RALES1999' -> 'RALES'; 'EMPEROR-Preserved' -> 'EMPERORPRESERVED'; 'PIONEER 6' -> 'PIONEER6'."""
    a = re.sub(r"(?<=[A-Za-z])((?:19|20)\d\d)$", "", _flat(a))
    return re.sub(r"[^A-Z0-9]", "", a.upper().replace("‐", "-"))


def aact_index(pmids, ncts, acronyms, agent_terms, root: str | None = None) -> dict:
    """ONE streaming pass per AACT table, restricted to the identifiers asked for.

    Returns pmid->ncts (study_references, any reference_type: a comparator citing a trial's PMID is a
    link claim we then check against RESULT/DERIVED), acronym->ncts (studies.acronym, INTERVENTIONAL
    only), per-NCT study facts, per-NCT intervention names, and per-NCT registered outcomes with
    results-posted flags. Fails closed (empty index + snapshot=None) when no snapshot is present."""
    from harness import aact
    snap = aact.snapshot_dir(root)
    idx = {"snapshot": snap, "pmid_nct": {}, "acr_nct": {}, "study": {}, "interventions": {}, "outcomes": {}}
    if not snap:
        return idx
    want_p = {str(p) for p in pmids if str(p).isdigit()}
    want_a = {norm_acronym(a) for a in acronyms if norm_acronym(a)}
    for r in aact._iter_rows(os.path.join(snap, "study_references.txt")):
        p = (r.get("pmid") or "").strip()
        if p in want_p:
            idx["pmid_nct"].setdefault(p, [])
            n = r["nct_id"].upper()
            if n not in [x[0] for x in idx["pmid_nct"][p]]:
                idx["pmid_nct"][p].append((n, (r.get("reference_type") or "").upper()))
    nct_set = {n.upper() for n in ncts} | {n for v in idx["pmid_nct"].values() for n, _ in v}
    for r in aact._iter_rows(os.path.join(snap, "studies.txt")):
        n = r["nct_id"].upper()
        acr = norm_acronym(r.get("acronym") or "")
        if acr and acr in want_a and (r.get("study_type") or "").upper() == "INTERVENTIONAL":
            idx["acr_nct"].setdefault(acr, []).append(n)
            nct_set.add(n)
    agent_re = re.compile("|".join(re.escape(a) for a in agent_terms), re.I) if agent_terms else None
    for r in aact._iter_rows(os.path.join(snap, "interventions.txt")):
        n = r["nct_id"].upper()
        if n in nct_set:
            idx["interventions"].setdefault(n, []).append(r.get("name") or "")
    for r in aact._iter_rows(os.path.join(snap, "studies.txt")):
        n = r["nct_id"].upper()
        if n in nct_set:
            idx["study"][n] = {k: r.get(k) or "" for k in (
                "acronym", "brief_title", "overall_status", "phase", "enrollment", "results_first_posted_date",
                "completion_date", "study_type")}
    for r in aact._iter_rows(os.path.join(snap, "outcomes.txt")):
        n = r["nct_id"].upper()
        if n in nct_set:
            idx["outcomes"].setdefault(n, []).append({k: (r.get(k) or "")[:300] for k in (
                "id", "outcome_type", "title", "time_frame", "population", "param_type", "units")})
    idx["agent_nct"] = {n: bool(agent_re and any(agent_re.search(x) for x in v))
                        for n, v in idx["interventions"].items()}
    return idx


# ------------------------------------------------------- model PROPOSAL of a comparator's members + the gate

MEMBERS_SCHEMA = {
    "type": "object", "additionalProperties": False, "required": ["state", "studies"],
    "properties": {
        "state": {"type": "string", "enum": ["ENUMERATED", "NOT_ENUMERATED"]},
        "studies": {"type": "array", "items": {
            "type": "object", "additionalProperties": False, "required": ["label", "quote", "design_stated"],
            "properties": {"label": {"type": "string"}, "quote": {"type": "string"},
                           "design_stated": {"type": "string", "enum": ["RCT", "NOT_RCT", "NOT_STATED"]}}}},
    },
}

MEMBERS_PROMPT = (
    "Below, between <<<TEXT and TEXT>>>, is the text of a published systematic review / meta-analysis.\n"
    "List EVERY individual study the review INCLUDED in its quantitative synthesis (any outcome). For each study give:\n"
    "- label: the study's name EXACTLY as the text writes it (a trial acronym such as 'RECOVERY', or a first-author "
    "label such as 'Meijvis 2011'). Copy it character for character.\n"
    "- quote: one verbatim passage (at most 300 characters) copied character for character from the text that "
    "contains the label and shows the study was included. Do not paraphrase, do not join passages, do not fix typos.\n"
    "- design_stated: RCT if the text says that study was randomised, NOT_RCT if it says otherwise, else NOT_STATED.\n"
    "Do not list excluded studies, ongoing studies, or studies only mentioned in the background/discussion. Do not use "
    "outside knowledge: a study the text does not name is not listed. If the text does not name its included studies, "
    "return state NOT_ENUMERATED with an empty list.\n<<<TEXT\n")


def _norm_ws(s: str) -> str:
    return _WS.sub(" ", (s or "").replace(" ", " ")).strip()


def verify_members(claim, held_text: str) -> dict:
    """Gate for a members PROPOSAL. Each study is admitted to the measurement only when its quote is in the
    WHOLE held text (whitespace-normalised on both sides -- the same bytes the model was shown) and its label
    is inside its own quote. A refused study is kept, with its reason, never dropped silently."""
    if not isinstance(claim, dict) or claim.get("state") not in ("ENUMERATED", "NOT_ENUMERATED"):
        return {"state": "VERIFIER_REFUSED", "problems": ["NOT_TYPED"], "admitted": [], "refused": []}
    held = _norm_ws(held_text)
    admitted, refused = [], []
    for s in claim.get("studies") or []:
        q, lab = _norm_ws(s.get("quote")), _norm_ws(s.get("label"))
        probs = []
        if not q or q not in held:
            probs.append("SPAN_NOT_IN_SOURCE")
        if not lab or lab not in q:
            probs.append("LABEL_NOT_IN_QUOTE")
        (refused if probs else admitted).append({**s, "problems": probs})
    return {"state": "VERIFIER_PASS" if admitted or claim["state"] == "NOT_ENUMERATED" else "VERIFIER_REFUSED",
            "claim_state": claim["state"], "admitted": admitted, "refused": refused}


def jats_body_text(body: bytes) -> str:
    """Plain text of a JATS article body + tables (the held text for a comparator with no citing table)."""
    root = ET.fromstring(body)
    b = root.find(".//body")
    parts = [_text(b)] if b is not None else []
    for tw in root.iter("table-wrap"):
        parts.append(_text(tw))
    return "\n".join(parts)


_STUDY_FIELDS = ("acronym", "brief_title", "overall_status", "phase", "enrollment", "results_first_posted_date",
                 "completion_date", "study_type", "study_first_submitted_date", "start_date")


class AactStore:
    """Persistent, incremental AACT index for repeated k-gap runs (the snapshot is read-only and slow to
    scan). Built maps: acronym -> NCTs (INTERVENTIONAL), pmid -> [(nct, RESULT|DERIVED)]; per-NCT facts
    (study, interventions, results outcomes) are scanned only for NCTs not already held. Keyed by the
    snapshot folder, so a newer snapshot starts a fresh store rather than mixing dates."""

    def __init__(self, path: str, root: str | None = None):
        from harness import aact
        self.aact = aact
        self.snap = aact.snapshot_dir(root)
        self.path = path
        d = {}
        if os.path.exists(path):
            with open(path, encoding="utf-8") as fh:
                d = json.load(fh)
        if d.get("snapshot") != self.snap:
            d = {"snapshot": self.snap}
        self.d = d
        for k in ("acr", "acr_title", "pmid", "study", "interventions", "outcomes", "design_groups"):
            self.d.setdefault(k, {})

    def save(self):
        with open(self.path, "w", encoding="utf-8") as fh:
            json.dump(self.d, fh)

    def _t(self, name):
        return os.path.join(self.snap, name + ".txt")

    def build_maps(self, log=print):
        if not self.snap or (self.d["acr"] and self.d["pmid"] and self.d["acr_title"]):
            return
        log("AACT: scanning studies (acronyms + parenthesised title acronyms)")
        self.d["acr"], self.d["acr_title"] = {}, {}
        for r in self.aact._iter_rows(self._t("studies")):
            if (r.get("study_type") or "").upper() != "INTERVENTIONAL":
                continue
            n = r["nct_id"].upper()
            a = norm_acronym(r.get("acronym") or "")
            if a:
                self.d["acr"].setdefault(a, []).append(n)
            # RE-LY (NCT00262600) and ROCKET AF (NCT00403767) carry an EMPTY acronym field; the acronym is
            # only in the title, parenthesised. A parenthesised token in a title is the registrant naming it.
            for t in (r.get("brief_title") or "", r.get("official_title") or ""):
                for m in _PAREN_ACRO.finditer(t):
                    ta = norm_acronym(m.group(1))
                    if len(ta) >= 4 and n not in self.d["acr_title"].get(ta, []):
                        self.d["acr_title"].setdefault(ta, []).append(n)
        log("AACT: scanning study_references")
        for r in self.aact._iter_rows(self._t("study_references")):
            t = (r.get("reference_type") or "").upper()
            p = (r.get("pmid") or "").strip()
            if p.isdigit() and t in ("RESULT", "DERIVED"):
                lst = self.d["pmid"].setdefault(p, [])
                n = r["nct_id"].upper()
                if [n, t] not in lst:
                    lst.append([n, t])
        self.save()

    def ensure_registration_dates(self, ncts, log=print):
        """study_first_submitted_date for held NCTs that lack it (one studies pass). A paper cannot be the RESULT of
        a trial registered after it was published -- the check that separates a trial's own report from later
        registrations that cite it."""
        want = {n.upper() for n in ncts if n and n.upper() in self.d["study"]
                and "study_first_submitted_date" not in self.d["study"][n.upper()]}
        if not want or not self.snap:
            return
        log(f"AACT: registration dates for {len(want)} NCTs")
        for r in self.aact._iter_rows(self._t("studies")):
            n = r["nct_id"].upper()
            if n in want:
                self.d["study"][n].update({k: r.get(k) or "" for k in _STUDY_FIELDS})
        for n in want:
            self.d["study"][n].setdefault("study_first_submitted_date", "")
        self.save()

    def ensure_ncts(self, ncts, log=print):
        want = {n.upper() for n in ncts if n} - set(self.d["study"])
        if not want or not self.snap:
            return
        log(f"AACT: facts for {len(want)} new NCTs")
        for n in want:
            self.d["study"][n] = {}
            self.d["interventions"][n] = []
            self.d["outcomes"][n] = []
        for r in self.aact._iter_rows(self._t("studies")):
            n = r["nct_id"].upper()
            if n in want:
                self.d["study"][n] = {k: r.get(k) or "" for k in _STUDY_FIELDS}
        for r in self.aact._iter_rows(self._t("interventions")):
            n = r["nct_id"].upper()
            if n in want:
                self.d["interventions"][n].append(r.get("name") or "")
        for r in self.aact._iter_rows(self._t("outcomes")):
            n = r["nct_id"].upper()
            if n in want:
                self.d["outcomes"][n].append({k: (r.get(k) or "")[:300] for k in (
                    "id", "outcome_type", "title", "time_frame", "population", "param_type", "units")})
        self.save()

    def ensure_design_groups(self, log=print):
        """design_groups (group_type + title) for every held NCT that lacks them: the arm TYPE is how a
        double-dummy active-comparator trial (placebo named in its intervention list) is told apart from a
        placebo-controlled one."""
        want = {n for n in self.d["study"] if n not in self.d["design_groups"]}
        if not want or not self.snap:
            return
        log(f"AACT: design_groups for {len(want)} NCTs")
        for n in want:
            self.d["design_groups"][n] = []
        for r in self.aact._iter_rows(self._t("design_groups")):
            n = r["nct_id"].upper()
            if n in want:
                self.d["design_groups"][n].append({"group_type": r.get("group_type") or "",
                                                   "title": (r.get("title") or "")[:200]})
        self.save()

    def index(self, agent_terms) -> dict:
        """The dict shape resolve_unit / aact_source consume, with agent flags for THIS topic's agents."""
        agent_re = re.compile("|".join(re.escape(a) for a in agent_terms), re.I) if agent_terms else None
        if not hasattr(self, "_rev"):
            self._rev = {}
            for p, v in self.d["pmid"].items():
                for n, _t in v:
                    self._rev.setdefault(n, []).append(p)
        return {"snapshot": self.snap, "acr_nct": self.d["acr"], "acr_title_nct": self.d["acr_title"],
                "nct_pmids": self._rev,
                "pmid_nct": {p: [tuple(x) for x in v] for p, v in self.d["pmid"].items()},
                "study": self.d["study"], "interventions": self.d["interventions"], "outcomes": self.d["outcomes"],
                "design_groups": self.d["design_groups"],
                "agent_nct": {n: bool(agent_re and any(agent_re.search(x) for x in v))
                              for n, v in self.d["interventions"].items()}}


def held_text(slug: str, date: str = "2026-09-28") -> tuple[str, str]:
    """(text, ref) of a topic's comparator as held: its own JATS body+tables when k_gap fetched open JATS,
    else the committed held document (records.json#comparator_fulltext or the document_ref file). ONE
    function for the proposal step and the table step, so the gate searches the bytes the model was shown."""
    with open(os.path.join(ROOT, "cache", slug, "comparators.json"), encoding="utf-8") as fh:
        c = json.load(fh)[0]
    m = re.search(r"PMID (\d+)", c.get("citation", ""))
    pmid = m.group(1) if m else str(c["id"])
    jp = os.path.join(COMP_DIR, pmid, f"{date}_kgap_jats.xml")
    if os.path.exists(jp):
        with open(jp, "rb") as fh:
            return jats_body_text(fh.read()), os.path.relpath(jp, ROOT).replace(os.sep, "/") + "#body"
    ref = c["document_ref"]
    with open(os.path.join(ROOT, ref), encoding="utf-8") as fh:
        raw = fh.read()
    if ref.endswith("records.json"):
        return json.loads(raw)["comparator_fulltext"], ref + "#comparator_fulltext"
    return raw, ref


# ------------------------------------------------- comparator REFERENCE SEEDING (open reference lists)

def fetch_epmc_references(pmid: str, date: str, offline: bool = False) -> dict:
    """The comparator's cited-reference list from Europe PMC (/MED/<pmid>/references), cached raw with
    sha256. Open metadata, so it exists for comparators whose full text is not open."""
    from harness import http
    name = f"{date}_kgap_epmc_refs.json"
    p = os.path.join(COMP_DIR, pmid, name)
    if os.path.exists(p):
        with open(p, "rb") as fh:
            body = fh.read()
    elif offline:
        return {"state": "NOT_CACHED_OFFLINE", "refs": []}
    else:
        pages, page = [], 1
        while page <= 5:
            try:
                d = http.get_json(f"https://www.ebi.ac.uk/europepmc/webservices/rest/MED/{pmid}/references",
                                  {"format": "json", "pageSize": 1000, "page": page})
            except Exception as exc:  # noqa: BLE001
                return {"state": "FETCH_FAILED", "error": str(exc)[:200], "refs": []}
            pages.append(d)
            if len(pages) * 1000 >= int(d.get("hitCount") or 0):
                break
            page += 1
        body = json.dumps(pages, sort_keys=True).encode("utf-8")
        _cache_write(pmid, name, body)
    pages = json.loads(body.decode("utf-8"))
    refs = []
    for d in pages:
        for r in (d.get("referenceList") or {}).get("reference", []):
            refs.append({"pmid": str(r.get("id") or "") if r.get("source") == "MED" else "",
                         "title": r.get("title") or "", "year": str(r.get("pubYear") or ""),
                         "author": (r.get("authorString") or "").split(",")[0], "doi": r.get("doi") or ""})
    return {"state": "REFS_HELD" if refs else "NO_REFS", "refs": refs, "sha256": sha256(body),
            "file": os.path.relpath(p, ROOT).replace(os.sep, "/")}


def pubmed_pubtypes(pmids, cache_path: str, offline: bool = False) -> dict:
    """PMID -> {pubtypes, title} via NCBI esummary (batched, cached)."""
    from harness import http
    cache = {}
    if os.path.exists(cache_path):
        with open(cache_path, encoding="utf-8") as fh:
            cache = json.load(fh)
    todo = [] if offline else sorted({p for p in pmids if p and p.isdigit() and p not in cache})
    for i in range(0, len(todo), 150):
        chunk = todo[i:i + 150]
        try:
            d = http.get_json("https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esummary.fcgi",
                              {"db": "pubmed", "id": ",".join(chunk), "retmode": "json"})
        except Exception as exc:  # noqa: BLE001
            print("esummary failed", exc)
            continue
        res = d.get("result", {})
        for p in chunk:
            r = res.get(p) or {}
            cache[p] = {"pubtypes": r.get("pubtype") or [], "title": r.get("title") or ""}
    if todo:
        with open(cache_path, "w", encoding="utf-8") as fh:
            json.dump(cache, fh, indent=1, sort_keys=True)
    return cache


def reference_seed_units(refs: list[dict], pubtypes: dict, agent_terms: list[str]) -> list[dict]:
    """Candidate members from a comparator's reference list: PubMed-typed RCT reports whose title names a
    topic agent. A SUPERSET candidate (a review cites trials it excludes), never the included set itself."""
    agent_re = re.compile("|".join(re.escape(a) for a in agent_terms), re.I) if agent_terms else None
    out = []
    for r in refs:
        pt = pubtypes.get(r["pmid"], {})
        title = pt.get("title") or r["title"]
        rct = any(t in ("Randomized Controlled Trial", "Clinical Trial, Phase III", "Clinical Trial, Phase II")
                  for t in pt.get("pubtypes", []))
        if r["pmid"] and rct and agent_re and agent_re.search(title):
            out.append({"table": "reference_seed", "layout": "reference", "label": title[:160], "context": title,
                        "rids": [], "cited": [{"pmid": r["pmid"], "doi": r.get("doi")}], "ncts": [],
                        "acronyms": [], "author": "", "year": r["year"], "agent_hit": True,
                        "drug_match": "DRUG_MATCH"})
    return out


def _shingles(t: str, n: int = 6) -> set:
    w = re.findall(r"[a-z0-9]+", (t or "").lower())
    return {" ".join(w[i:i + n]) for i in range(len(w) - n + 1)}


def held_text_identity(abstract: str, held: str) -> dict:
    """Is the held comparator text the article whose PubMed abstract this is? Measured by shared 6-word
    shingles. On the 32 served comparators (2026-09-28) the correct articles share 22-874 abstract shingles;
    two held texts share ZERO -- they are different articles. A bag-of-title-words check was tried first and
    passed both (9/10 and 6/7 title words present): topical vocabulary is not identity."""
    a, h = _shingles(abstract), _shingles(held)
    shared = len(a & h)
    state = ("NO_ABSTRACT" if not a else "NAMED_ARTICLE" if shared >= 10 else
             "HELD_TEXT_NOT_NAMED_ARTICLE" if shared == 0 else "IDENTITY_UNCERTAIN")
    return {"abstract_shingles": len(a), "shared": shared, "state": state}


_REFNUM = re.compile(r"[\[(]\s*(\d{1,3})\s*[\])]|\^(\d{1,3})|(?<=[A-Za-z])(\d{1,3})$")


def refs_by_number(label: str, refs: dict) -> list[dict]:
    """A proposal label copied with its citation number ('Imazio [19]', 'Zinman (8)') cites the comparator's
    own reference list. Resolve the number by the ref's JATS <label> first, else by its ordinal position; the
    two agreeing is not required, but a number matching NEITHER resolves nothing (never a nearest guess)."""
    nums = [int(next(g for g in m.groups() if g)) for m in _REFNUM.finditer(label or "")]
    out = []
    for n in nums:
        hit = [r for r in refs.values() if r.get("label") == str(n)] or \
              [r for r in refs.values() if r.get("ordinal") == n and not r.get("label")]
        if len(hit) == 1:
            r = hit[0]
            sur = re.match(r"^\s*([A-Z][A-Za-z'À-ſ-]{1,})", label or "")
            if sur and r.get("first_author") and not (
                    r["first_author"].lower().startswith(sur.group(1).lower()[:5])
                    or sur.group(1).lower().startswith(r["first_author"].lower()[:5])):
                continue            # the surname on the label and the cited ref disagree: resolve nothing
            out.append(r)
    return out


# ---------------------------------------------------------------------------- Unpaywall OA copy -> text

UPW_CONTACT = "meta-harness@example.org"      # the harness's own contact; never a person's address


def _pdf_text(data: bytes) -> str:
    try:
        import io as _io
        from pypdf import PdfReader
        return "\n".join((p.extract_text() or "") for p in PdfReader(_io.BytesIO(data)).pages)
    except Exception:  # noqa: BLE001
        return ""


def _html_text(markup: str) -> str:
    t = re.sub(r"(?is)<(script|style|noscript|svg).*?</\1>", " ", markup or "")
    return _flat(re.sub(r"(?s)<[^>]+>", " ", t))


def unpaywall_text(doi: str, cache_dir: str, index_path: str, offline: bool = False) -> dict:
    """Legitimate OA copy of a DOI via Unpaywall (best + other OA locations), as TEXT (PDF via pypdf, else HTML).
    Only the text is kept (cache_dir, gitignored); index_path records DOI -> source url, kind, sha256, bytes.
    A publisher PDF has no table delimiters, so baseline tables cannot be told from outcome tables in it: callers
    must treat Unpaywall text as prose-only evidence and audit every admission from it."""
    from harness import http
    import urllib.parse
    os.makedirs(cache_dir, exist_ok=True)
    key = hashlib.sha1(doi.lower().encode("utf-8")).hexdigest()[:16]
    fp = os.path.join(cache_dir, key + ".txt")
    idx = {}
    if os.path.exists(index_path):
        with open(index_path, encoding="utf-8") as fh:
            idx = json.load(fh)
    if os.path.exists(fp):
        with open(fp, encoding="utf-8") as fh:
            return {"text": fh.read(), **idx.get(doi.lower(), {})}
    if offline:
        return {"text": "", "state": "NOT_CACHED_OFFLINE"}
    out = {"state": "NO_OA_TEXT", "tried": []}
    try:
        d = http.get_json("https://api.unpaywall.org/v2/" + urllib.parse.quote(doi, safe=""),
                          {"email": UPW_CONTACT}, tries=2)
    except Exception as exc:  # noqa: BLE001
        out["state"] = "UNPAYWALL_ERROR"
        out["error"] = str(exc)[:160]
        d = {}
    locs = ([d.get("best_oa_location")] if d.get("best_oa_location") else []) + list(d.get("oa_locations") or [])
    seen, text = set(), ""
    for loc in locs:
        for k in ("url_for_pdf", "url"):
            u = (loc or {}).get(k)
            if not u or u in seen:
                continue
            seen.add(u)
            try:
                st, b = http.get_raw(u, tries=2, timeout=45)
            except Exception as exc:  # noqa: BLE001
                out["tried"].append({"url": u, "error": str(exc)[:100]})
                continue
            kind = "PDF" if b[:4] == b"%PDF" else "HTML"
            t = _pdf_text(b) if kind == "PDF" else _html_text(b.decode("utf-8", "replace"))
            out["tried"].append({"url": u, "kind": kind, "http_status": st, "text_bytes": len(t)})
            if len(t) >= 3000:
                text = t
                out.update({"state": "OA_TEXT", "url": u, "kind": kind, "host_type": (loc or {}).get("host_type"),
                            "license": (loc or {}).get("license")})
                break
        if text:
            break
    with open(fp, "w", encoding="utf-8") as fh:
        fh.write(text)
    tb = text.encode("utf-8")
    # keep WHY each location gave no text (status, kind, text bytes, error): dropping 'tried' left 65 "OA copy
    # found" DOIs with no recorded reason for yielding nothing
    idx[doi.lower()] = {k: v for k, v in out.items() if k != "tried"} | {"sha256": sha256(tb), "bytes": len(tb),
                                                                          "n_tried": len(out["tried"]),
                                                                          "tried": out["tried"],
                                                                          "n_oa_locations": len(locs)}
    with open(index_path, "w", encoding="utf-8") as fh:
        json.dump(idx, fh, indent=1, sort_keys=True)
    return {"text": text, **idx[doi.lower()]}


_GENERIC_LABEL_WORDS = {"trial", "trials", "study", "studies", "cohort", "group", "arm", "patients", "participants",
                        "author", "authors", "reference", "ref", "the", "and", "phase", "part", "sub", "substudy"}


def label_ref_conflict(unit: dict, cited: dict, registry_acronyms=()) -> str | None:
    """Does a table row's own label contradict the reference its citation link points to? The omega-3 comparator
    (PMID 35905212, PMC JATS) cites 'GISSI-HF 2008 [33]' where ref 33 is JELIS (Yokoyama 2007), and 'Kromhout 2010
    [36]' where ref 36 is a DHA-in-Alzheimer trial (Quinn 2010): the publisher's numbering is off, so following the
    link faithfully resolves the WRONG trial. Returns the conflict (author / acronym / year) or None when consistent
    or when the label carries nothing to compare.

    Only CONTRADICTION counts, never absence: an article title routinely omits the trial acronym (HEART-FID, STEP 1),
    so an acronym is judged only against the cited paper's REGISTRY acronyms (registry_acronyms, from AACT), and only
    when there are some. The cited ref's own number is stripped from the label first ('STEP 1 29' cites ref 29; the
    acronym is 'STEP 1'), accents are folded (Garzon == Garzon-with-acute), and a caps-heavy first token
    ('SCALEObesity', 'SURMOUNT-') is an acronym, not a surname."""
    import unicodedata

    def fold(s):
        return unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode("ascii").lower()

    label = unit.get("label") or ""
    num = (cited.get("label") or "").strip()
    if num.isdigit():
        label = re.sub(rf"(?:[\[(\s]|(?<=[A-Za-z]))\s*{num}\s*[\])]?\s*$", "", label).strip()
    au, yr = unit.get("author") or "", unit.get("year") or ""
    if not au:
        m = re.match(r"^\s*([A-Z][A-Za-z'À-ſ-]{2,})\b", label)
        au = m.group(1) if m else ""
    if au and (_ACRO.fullmatch(au) or len(re.findall(r"[A-Z]", au)) >= 2 or au.endswith("-")):
        au = ""                                    # an acronym-shaped token is not a surname
    if au.lower() in _GENERIC_LABEL_WORDS:
        au = ""                                    # 'Trial A (2019)' (esketamine comparator) names no author
    if not yr:
        m = re.search(r"\b((?:19|20)\d\d)\b", label)
        yr = m.group(1) if m else ""
    fa = fold(cited.get("first_author"))
    if au and fa and fa not in ("group", "investigators", "writing", "collaborators", "committee") and not (
            fa.startswith(fold(au)[:5]) or fold(au).startswith(fa[:5])):
        return f"author: row '{au}' vs cited first author '{cited.get('first_author')}'"
    # the cited paper naming itself counts too: '(JELIS)' in its title contradicts a 'GISSI-HF' row
    named = [m.group(1) for m in _PAREN_ACRO.finditer(cited.get("title") or "")]
    reg = [norm_acronym(a) for a in list(registry_acronyms or ()) + named if len(norm_acronym(a)) >= 4]
    acr = [norm_acronym(re.sub(r"\s+\d{1,3}$", "", a)) for a in (unit.get("acronyms") or [])]
    acr = [a for a in acr if len(a) >= 4]
    if acr and reg and not any(a in r or r in a for a in acr for r in reg):
        return f"acronym: row '{acr[0]}' vs cited registry {reg[:3]}"
    ry = cited.get("year") or ""
    if yr and ry and ry.isdigit() and abs(int(yr) - int(ry)) > 1:
        return f"year: row {yr} vs cited {ry}"
    return None


# ------------------------------------------------------------ comparator SUPPLEMENTS (PMC OA package) -> text

def _docx_text(data: bytes) -> str:
    """Paragraphs, then every table row as 'cell | cell' (verbatim cell text)."""
    try:
        import io as _io
        import docx
        d = docx.Document(_io.BytesIO(data))
        out = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows:
                cells = []
                for c in r.cells:
                    tx = _flat(c.text)
                    if not cells or cells[-1] != tx:        # merged cells repeat; keep one
                        cells.append(tx)
                out.append(" | ".join(cells))
        return "\n".join(out)
    except Exception:  # noqa: BLE001
        return ""


def comparator_supplements(pmid: str, pmcid: str, hrefs: list, date: str, offline: bool = False) -> dict:
    """Text of a comparator's supplementary files from its PMC OA package (docx: paragraphs + table rows; pdf: pypdf
    text; xlsx/csv: the harness's own readers). Figures/images are skipped (no OCR). Cached as text with sha256 under
    cache/comparators/<pmid>/<date>_kgap_supplements.txt. Returns {state, text, files}."""
    from harness import fulltext as _ft
    from harness import http
    fp = os.path.join(COMP_DIR, pmid, f"{date}_kgap_supplements.txt")
    mp = fp[:-4] + ".manifest.json"
    if os.path.exists(fp):
        with open(fp, encoding="utf-8") as fh:
            t = fh.read()
        side = {}
        if os.path.exists(mp):
            with open(mp, encoding="utf-8") as fh:
                side = json.load(fh)
        if side.get("sha256") and side["sha256"] != sha256(t.encode("utf-8")):
            return {"state": "CACHE_HASH_MISMATCH", "text": "", "route": side.get("route")}
        return {**side, "state": "CACHED", "fetched_state": side.get("state"), "text": t,
                "sha256": sha256(t.encode("utf-8"))}
    if offline:
        return {"state": "NOT_CACHED_OFFLINE", "text": ""}
    if not pmcid or not hrefs:
        return {"state": "NO_SUPPLEMENTS", "text": ""}
    # Routes, in order. NCBI's oa.fcgi answers 404 (retired; harness.fetch._pmc_oa_supplement_text swallows that and
    # returns ''), and PMC's articles/instance/<id>/bin/<file> serves a JavaScript interstitial (bot protection; not
    # circumvented). Europe PMC's REST supplementaryFiles endpoint serves every supplement of an OA article as one ZIP.
    url = f"https://www.ebi.ac.uk/europepmc/webservices/rest/{pmcid}/supplementaryFiles"
    try:
        st, zbytes = http.get_raw(url, tries=3, timeout=180)
    except Exception as exc:  # noqa: BLE001
        return {"state": "FETCH_FAILED", "text": "", "route": url, "error": str(exc)[-200:]}
    if zbytes[:2] != b"PK":          # an XML/HTML error body is not a package: fail closed, cache nothing
        return {"state": "NOT_A_ZIP", "text": "", "route": url, "http_status": st, "bytes": len(zbytes),
                "head": zbytes[:120].decode("utf-8", "replace")}
    wanted = {h.rsplit("/", 1)[-1].lower() for h in hrefs}
    blocks, files = [], []
    for base, data in _zip_members(zbytes):
        low = base.lower()
        rec = {"file": base, "bytes": len(data), "sha256": sha256(data), "listed": low in wanted}
        if low.endswith((".jpg", ".jpeg", ".png", ".gif", ".tif", ".tiff", ".mp4", ".bmp")):
            files.append({**rec, "skipped": "image (no OCR)"})
            continue
        if low.endswith(".docx"):
            txt = _docx_text(data)
        elif low.endswith(".pdf") or data[:4] == b"%PDF":
            txt = _pdf_text(data)
        elif low.endswith(".doc"):
            files.append({**rec, "skipped": "legacy binary .doc (no typed parser)"})
            continue
        else:
            txt = _ft.supplement_text_from_bytes(low, data)
        files.append({**rec, "text_chars": len(txt or "")})
        if txt:
            blocks.append(f"=== SUPPLEMENT {base} ===\n{txt}")
    text = "\n\n".join(blocks)
    os.makedirs(os.path.dirname(fp), exist_ok=True)
    with open(fp, "w", encoding="utf-8") as fh:
        fh.write(text)
    side = {"state": "OK" if text else "NO_TEXT_IN_SUPPLEMENTS", "files": files, "route": url, "http_status": st,
            "zip_bytes": len(zbytes), "zip_sha256": sha256(zbytes), "sha256": sha256(text.encode("utf-8"))}
    with open(mp, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(side, fh, indent=1)
    return {**side, "text": text}


def _zip_members(zbytes: bytes, depth: int = 0):
    """(basename, bytes) for every file in a ZIP, descending into nested ZIPs (MDPI ships one) up to depth 2."""
    import io as _io
    import zipfile
    with zipfile.ZipFile(_io.BytesIO(zbytes)) as z:
        for info in z.infolist():
            if info.is_dir():
                continue
            data = z.read(info)
            base = info.filename.rsplit("/", 1)[-1]
            if base.lower().endswith(".zip") and depth < 2 and data[:2] == b"PK":
                yield from _zip_members(data, depth + 1)
            else:
                yield base, data
